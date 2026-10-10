"""Вкладка проєктів (PySide6 + PostgreSQL) — v2.4 зі знижкою на виробах.

Логіка:
  Собівартість = сума cost_price × quantity з усіх виробів  [auto]
  Ціна замовнику = сума discounted_price (якщо > 0) інакше total_price  [auto]
  Ціна зі знижкою (проєкт) = ви вводите самі                [editable]
  Прибуток = ціна (зі знижкою) − собівартість               [auto]
"""

import contextlib
import csv
from datetime import datetime

from PySide6.QtGui import QBrush, QColor, QStandardItem, QStandardItemModel
from PySide6.QtWidgets import (
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QDoubleSpinBox,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QTableView,
    QVBoxLayout,
    QWidget,
)

from ventilation_company.database.repositories.client_repo import ClientRepository
from ventilation_company.database.repositories.product_repo import ProductRepository
from ventilation_company.database.repositories.project_repo import ProjectRepository
from ventilation_company.gui_pyside6.project_card_dialog import ProjectCardDialog
from ventilation_company.gui_pyside6.table_utils import setup_table
from ventilation_company.gui_pyside6.theme import Theme
from ventilation_company.gui_pyside6.workers import FunctionWorker
from ventilation_company.services.audit_service import log_action
from ventilation_company.services.project_notifications import (
    notify_project_created,
    notify_project_status_changed,
)
from ventilation_company.utils.logging_config import get_logger


class ProjectEditDialog(QDialog):
    """Діалог створення/редагування проєкту з ціноутворенням з виробів."""

    STATUSES = ["Новий", "В роботі", "На виробництві", "Готовий", "Відвантажено", "Закрито"]

    def __init__(self, project_data=None, parent=None):
        super().__init__(parent)
        self.project_data = project_data or {}
        self.setWindowTitle("Редагування проєкту" if project_data else "Новий проєкт")
        self.setMinimumWidth(480)
        self._build_ui()
        self._recalc_profit()

    def _build_ui(self):
        layout = QFormLayout(self)
        layout.setSpacing(10)

        project_id = self.project_data.get("id")
        products = []
        if project_id:
            with contextlib.suppress(Exception):
                products = ProductRepository.get_all(project_id=project_id)

        cost_total = sum(p.get("cost_price", 0) * p.get("quantity", 1) for p in products)
        # ← v2.4: ціна з урахуванням знижки на вироби
        price_total = sum(
            (
                p.get("discounted_price", 0)
                if p.get("discounted_price", 0) > 0
                else p.get("total_price", 0)
            )
            for p in products
        )

        self.edit_name = QLineEdit()
        self.edit_name.setText(self.project_data.get("name", ""))
        layout.addRow("Назва *", self.edit_name)

        self.edit_number = QLineEdit()
        self.edit_number.setText(self.project_data.get("project_number", ""))
        layout.addRow("Номер", self.edit_number)

        self.edit_contract = QLineEdit()
        self.edit_contract.setText(self.project_data.get("contract_number") or "")
        self.edit_contract.setPlaceholderText("Автоматично при формуванні договору")
        layout.addRow("Договір", self.edit_contract)

        self.edit_invoice = QLineEdit()
        self.edit_invoice.setText(self.project_data.get("invoice_number") or "")
        self.edit_invoice.setPlaceholderText("Автоматично при формуванні рахунку")
        layout.addRow("Рахунок", self.edit_invoice)

        self.edit_act = QLineEdit()
        self.edit_act.setText(self.project_data.get("act_number") or "")
        self.edit_act.setPlaceholderText("Автоматично при формуванні акту")
        layout.addRow("Акт", self.edit_act)

        self.combo_client = QComboBox()
        self.combo_client.setEditable(True)
        try:
            clients = ClientRepository.list_all()
        except Exception:
            clients = []
        for client in clients:
            label = client.get("name") or f"Клієнт #{client['id']}"
            extra = client.get("phone") or client.get("email") or ""
            display = f"{label} ({extra})" if extra else label
            self.combo_client.addItem(display, client["id"])
        current_client_id = self.project_data.get("client_id")
        if current_client_id:
            idx = self.combo_client.findData(current_client_id)
            if idx >= 0:
                self.combo_client.setCurrentIndex(idx)
        elif self.project_data.get("client"):
            self.combo_client.setCurrentText(str(self.project_data.get("client")))
        layout.addRow("Клієнт", self.combo_client)

        self.combo_status = QComboBox()
        self.combo_status.addItems(self.STATUSES)
        current = self.project_data.get("status", "Новий")
        idx = self.combo_status.findText(current)
        if idx >= 0:
            self.combo_status.setCurrentIndex(idx)
        layout.addRow("Статус", self.combo_status)

        layout.addRow(QLabel(""))

        self.spin_cost = QDoubleSpinBox()
        self.spin_cost.setRange(0, 99999999)
        self.spin_cost.setSuffix(" ₴")
        self.spin_cost.setDecimals(2)
        self.spin_cost.setValue(cost_total)
        self.spin_cost.setReadOnly(True)
        self.spin_cost.setButtonSymbols(QDoubleSpinBox.ButtonSymbols.NoButtons)
        self.spin_cost.setStyleSheet(
            "QDoubleSpinBox { background-color: #1a1a2e; color: #a6e3a1; font-weight: bold; }"
        )
        layout.addRow("🔧 Собівартість (з виробів)", self.spin_cost)

        self.spin_base_price = QDoubleSpinBox()
        self.spin_base_price.setRange(0, 99999999)
        self.spin_base_price.setSuffix(" ₴")
        self.spin_base_price.setDecimals(2)
        self.spin_base_price.setValue(price_total)
        self.spin_base_price.setReadOnly(True)
        self.spin_base_price.setButtonSymbols(QDoubleSpinBox.ButtonSymbols.NoButtons)
        self.spin_base_price.setStyleSheet(
            "QDoubleSpinBox { background-color: #1a1a2e; color: #89b4fa; font-weight: bold; }"
        )
        layout.addRow("💰 Ціна замовнику (з виробів)", self.spin_base_price)

        self.spin_discounted = QDoubleSpinBox()
        self.spin_discounted.setRange(0, 99999999)
        self.spin_discounted.setSuffix(" ₴")
        self.spin_discounted.setDecimals(2)
        self.spin_discounted.setValue(float(self.project_data.get("discounted_price", 0) or 0))
        self.spin_discounted.setStyleSheet(
            "QDoubleSpinBox { background-color: #2a2a3e; color: #f9e2af; font-weight: bold; }"
        )
        self.spin_discounted.valueChanged.connect(self._recalc_profit)
        layout.addRow("🏷️ Ціна зі знижкою", self.spin_discounted)

        self.lbl_profit = QLabel("0.00 ₴")
        self.lbl_profit.setStyleSheet(
            "font-size: 14px; font-weight: bold; padding: 6px 12px; border-radius: 6px; background: #1a1a2e;"
        )
        layout.addRow("📊 Прибуток", self.lbl_profit)

        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addRow(buttons)

    def _recalc_profit(self):
        cost = self.spin_cost.value()
        base_price = self.spin_base_price.value()
        discounted = self.spin_discounted.value()
        effective_price = discounted if discounted > 0 else base_price
        profit = effective_price - cost
        if profit >= 0:
            self.lbl_profit.setText(f"{profit:,.2f} ₴  ✅")
            self.lbl_profit.setStyleSheet(
                "font-size: 14px; font-weight: bold; padding: 6px 12px; "
                "border-radius: 6px; background: #1a3a1a; color: #a6e3a1;"
            )
        else:
            self.lbl_profit.setText(f"{profit:,.2f} ₴  ⚠️ ЗБИТОК")
            self.lbl_profit.setStyleSheet(
                "font-size: 14px; font-weight: bold; padding: 6px 12px; "
                "border-radius: 6px; background: #3a1a1a; color: #f38ba8;"
            )

    def get_data(self):
        cost = self.spin_cost.value()
        base_price = self.spin_base_price.value()
        discounted = self.spin_discounted.value()
        effective_price = discounted if discounted > 0 else base_price
        profit = effective_price - cost
        return {
            "name": self.edit_name.text().strip(),
            "project_number": self.edit_number.text().strip(),
            "contract_number": self.edit_contract.text().strip(),
            "invoice_number": self.edit_invoice.text().strip(),
            "act_number": self.edit_act.text().strip(),
            "client": self.combo_client.currentText().strip(),
            "client_id": self.combo_client.currentData(),
            "status": self.combo_status.currentText(),
            "cost_price": cost,
            "customer_price": base_price,
            "discounted_price": discounted,
            "profit": profit,
        }


class ProjectsTab(QWidget):
    def __init__(self, parent=None, main_window=None):
        super().__init__(parent)
        self.main_window = main_window
        self._projects = []
        self._build_ui()
        self._load_data()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.setSpacing(16)

        header = QHBoxLayout()
        buttons_row = QHBoxLayout()
        buttons_row.setSpacing(6)
        lbl_title = QLabel("📁 Проєкти")
        lbl_title.setObjectName("title")
        header.addWidget(lbl_title)
        header.addStretch()

        self.edit_search = QLineEdit()
        self.edit_search.setMaximumWidth(170)
        self.edit_search.setPlaceholderText("🔍 Пошук проєкту...")
        self.edit_search.setFixedWidth(250)
        self.edit_search.textChanged.connect(self._on_search)
        self.combo_status_filter = QComboBox()
        self.combo_status_filter.setMaximumWidth(120)
        self.combo_status_filter.addItems(
            ["Всі", "Новий", "В роботі", "На виробництві", "Готовий", "Відвантажено", "Закрито"]
        )
        self.combo_status_filter.currentTextChanged.connect(
            lambda _text: self._on_search(self.edit_search.text())
        )
        header.addWidget(self.edit_search)
        header.addWidget(self.combo_status_filter)

        btn_refresh = QPushButton("🔄 Оновити")
        btn_export_csv = QPushButton("💾 CSV")
        btn_export_csv.clicked.connect(self._export_csv)
        btn_export_excel = QPushButton("📊 Excel")
        btn_export_excel.setToolTip("Експорт списку проєктів у Excel з сумами та прибутком")
        btn_export_excel.clicked.connect(self._export_excel)
        btn_refresh.clicked.connect(self._load_data)
        buttons_row.addWidget(btn_refresh)
        buttons_row.addWidget(btn_export_csv)
        buttons_row.addWidget(btn_export_excel)

        btn_new = QPushButton("➕ Новий")
        btn_new.setObjectName("primary")
        btn_new.clicked.connect(self._on_new_project)
        btn_duplicate = QPushButton("📄 Дублювати")
        btn_status_work = QPushButton("▶ Робота")
        btn_status_ready = QPushButton("✅ Готово")
        btn_status_closed = QPushButton("🏁 Закрито")
        btn_status_work.clicked.connect(lambda: self._set_selected_status("В роботі"))
        btn_status_ready.clicked.connect(lambda: self._set_selected_status("Готовий"))
        btn_status_closed.clicked.connect(lambda: self._set_selected_status("Закрито"))
        btn_duplicate.clicked.connect(self._on_duplicate_project)
        buttons_row.addWidget(btn_new)
        buttons_row.addWidget(btn_duplicate)
        buttons_row.addWidget(btn_status_work)
        buttons_row.addWidget(btn_status_ready)
        buttons_row.addWidget(btn_status_closed)

        layout.addLayout(header)
        layout.addLayout(buttons_row)

        self.table = QTableView()
        setup_table(
            self.table,
            select_rows=True,
            single_selection=True,
            sorting=True,
            read_only=True,
        )
        self.table.setMinimumHeight(400)
        # activated = Enter або подвійний клік → картка проєкту
        self.table.activated.connect(self._on_double_click)
        layout.addWidget(self.table)

        self.model = QStandardItemModel()
        self.model.setHorizontalHeaderLabels(
            [
                "ID",
                "Номер",
                "Назва",
                "Клієнт",
                "Статус",
                "Договір",
                "Дата створення",
                "Сума виробів",
                "Ціна зам.",
                "Зі знижкою",
                "Прибуток",
            ]
        )
        self.table.setModel(self.model)
        self.table.setColumnWidth(0, 40)
        self.table.setColumnWidth(1, 90)
        self.table.setColumnWidth(2, 180)
        self.table.setColumnWidth(3, 130)
        self.table.setColumnWidth(4, 90)
        self.table.setColumnWidth(5, 120)
        self.table.setColumnWidth(6, 90)
        self.table.setColumnWidth(7, 90)
        self.table.setColumnWidth(8, 90)
        self.table.setColumnWidth(9, 90)
        self.table.setColumnWidth(10, 90)

        actions = QHBoxLayout()
        actions.addStretch()
        btn_edit = QPushButton("✏️ Редагувати")
        btn_edit.clicked.connect(self._on_edit)
        actions.addWidget(btn_edit)
        btn_del = QPushButton("🗑️ Видалити")
        btn_del.setStyleSheet(f"color: {Theme.DANGER};")
        btn_del.clicked.connect(self._on_delete)
        actions.addWidget(btn_del)
        layout.addLayout(actions)

        lbl_hint = QLabel("💡 Двічі клікніть на рядок для відкриття картки проєкту")
        lbl_hint.setObjectName("subtitle")
        layout.addWidget(lbl_hint)

    def _get_selected_id(self):
        idx = self.table.currentIndex()
        if not idx.isValid():
            return None
        row = idx.row()
        if 0 <= row < len(self._projects):
            return self._projects[row].get("id")
        return None

    def _export_csv(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Експорт проєктів", "projects.csv", "CSV (*.csv)"
        )
        if not path:
            return
        fieldnames = [
            "id",
            "project_number",
            "name",
            "client",
            "status",
            "contract_number",
            "created_at",
            "cost_price",
            "customer_price",
            "discounted_price",
            "profit",
        ]
        rows = getattr(self, "_projects", [])
        try:
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                for row in rows:
                    writer.writerow({k: row.get(k, "") for k in fieldnames})
            QMessageBox.information(self, "Успіх", f"Експортовано проєктів: {len(rows)}")
        except Exception as e:
            QMessageBox.critical(self, "Помилка", f"Не вдалося експортувати CSV: {e}")

    def _export_excel(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Експорт проєктів у Excel", "projects.xlsx", "Excel (*.xlsx)"
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        rows = getattr(self, "_projects", [])
        headers = [
            "ID",
            "Номер",
            "Назва",
            "Клієнт",
            "Статус",
            "Договір",
            "Дата створення",
            "Собівартість",
            "Ціна замовника",
            "Зі знижкою",
            "Прибуток",
        ]
        keys = [
            "id",
            "project_number",
            "name",
            "client",
            "status",
            "contract_number",
            "created_at",
            "cost_price",
            "customer_price",
            "discounted_price",
            "profit",
        ]
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = "Проєкти"
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
            for col, title in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=title)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            money_cols = {8, 9, 10, 11}
            for r, row in enumerate(rows, 2):
                for col, key in enumerate(keys, 1):
                    value = row.get(key, "")
                    cell = ws.cell(row=r, column=col, value=value)
                    if col in money_cols:
                        cell.number_format = "#,##0.00"
                        cell.alignment = Alignment(horizontal="right")
            widths = [6, 14, 30, 24, 12, 16, 14, 14, 16, 12, 14]
            for col, width in enumerate(widths, 1):
                ws.column_dimensions[chr(64 + col)].width = width
            ws.freeze_panes = "A2"
            wb.save(path)
            QMessageBox.information(self, "Успіх", f"Експортовано проєктів: {len(rows)}")
        except Exception as e:
            QMessageBox.critical(self, "Помилка", f"Не вдалося експортувати Excel: {e}")

    def _load_data(self):
        self.model.removeRows(0, self.model.rowCount())
        self._projects = []
        try:
            projects = ProjectRepository.list_all()
            for p in projects:
                try:
                    products = ProductRepository.get_all(project_id=p["id"])
                    cost = sum(pr.get("cost_price", 0) * pr.get("quantity", 1) for pr in products)
                    base = sum(
                        (
                            pr.get("discounted_price", 0)
                            if pr.get("discounted_price", 0) > 0
                            else pr.get("total_price", 0)
                        )
                        for pr in products
                    )
                except Exception:
                    cost = 0
                    base = 0

                discounted = float(p.get("discounted_price") or 0)
                effective = discounted if discounted > 0 else base
                profit = effective - cost
                data = {
                    "id": p["id"],
                    "name": p.get("name") or "—",
                    "project_number": p.get("project_number") or "—",
                    "client": p.get("client") or "—",
                    "status": p.get("status") or "Новий",
                    "contract_number": p.get("contract_number") or "",
                    "created_at": str(p.get("created_at"))[:10] if p.get("created_at") else "—",
                    "cost_price": cost,
                    "customer_price": base,
                    "discounted_price": discounted,
                    "profit": profit,
                }
                self._projects.append(data)

                profit_color = Theme.SUCCESS if profit >= 0 else Theme.DANGER
                row = [
                    QStandardItem(str(p["id"])),
                    QStandardItem(data["project_number"]),
                    QStandardItem(data["name"]),
                    QStandardItem(data["client"]),
                    QStandardItem(data["status"]),
                    QStandardItem(data["contract_number"] or "—"),
                    QStandardItem(data["created_at"]),
                    QStandardItem("₴ " + f"{cost:,.0f}"),
                    QStandardItem("₴ " + f"{base:,.0f}"),
                    QStandardItem("₴ " + f"{discounted:,.0f}"),
                    QStandardItem("₴ " + f"{profit:,.0f}"),
                ]
                for cell in row:
                    cell.setEditable(False)
                row[-1].setForeground(QBrush(QColor(profit_color)))
                self.model.appendRow(row)
        except Exception as e:
            QMessageBox.critical(self, "Помилка", f"Не вдалося завантажити проєкти: {e}")

    def _on_search(self, text):
        text = text.lower()
        status_filter = (
            self.combo_status_filter.currentText()
            if hasattr(self, "combo_status_filter")
            else "Всі"
        )
        for row in range(self.model.rowCount()):
            visible = False
            for col in range(self.model.columnCount()):
                item = self.model.item(row, col)
                if item and text in item.text().lower():
                    visible = True
                    break
            if visible and status_filter != "Всі":
                status_item = self.model.item(row, 4)
                visible = bool(status_item and status_item.text() == status_filter)
            self.table.setRowHidden(row, not visible)

    def _on_double_click(self, index):
        row = index.row()
        if 0 <= row < len(self._projects):
            project_id = self._projects[row]["id"]
            dlg = ProjectCardDialog(project_id, parent=self)
            dlg.exec()

    def _audit_actor(self):
        return getattr(self.main_window, "user", None)

    def _set_selected_status(self, status: str):
        project_id = self._get_selected_id()
        if not project_id:
            QMessageBox.warning(self, "Увага", "Виберіть проєкт")
            return
        try:
            ProjectRepository.update(project_id, {"status": status})
            log_action(
                "project.update",
                entity_type="project",
                entity_id=project_id,
                details={"status": status},
                actor=self._audit_actor(),
            )
            self._load_data()
        except Exception as e:
            QMessageBox.critical(self, "Помилка", f"Не вдалося змінити статус: {e}")

    def _on_duplicate_project(self):
        project_id = self._get_selected_id()
        if not project_id:
            QMessageBox.warning(self, "Увага", "Виберіть проєкт для дублювання")
            return
        source = ProjectRepository.get(project_id)
        if not source:
            QMessageBox.warning(self, "Увага", "Проєкт не знайдено")
            return
        products = ProductRepository.get_all(project_id=project_id)
        answer = QMessageBox.question(
            self,
            "Дублювання проєкту",
            f"Створити копію проєкту '{source.get('name')}' з {len(products)} виробами?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return
        data = {
            "name": f"{source.get('name') or 'Проєкт'} (копія)",
            "project_number": f"PRJ-{datetime.now().strftime('%Y%m%d-%H%M%S')}",
            "client": source.get("client"),
            "address": source.get("address"),
            "status": "Новий",
            "start_date": source.get("start_date"),
            "deadline": source.get("deadline"),
            "notes": source.get("notes"),
            "created_at": datetime.now(),
        }
        try:
            new_project = ProjectRepository.create(data)
            new_id = new_project["id"]
            copied = 0
            for item in products:
                pdata = {
                    "name": item.get("name"),
                    "product_type": item.get("product_type"),
                    "width": item.get("width"),
                    "height": item.get("height"),
                    "length": item.get("length"),
                    "thickness": item.get("thickness"),
                    "material": item.get("material"),
                    "quantity": item.get("quantity"),
                    "cost_price": item.get("cost_price"),
                    "unit_price": item.get("unit_price"),
                    "total_price": item.get("total_price"),
                    "discounted_price": item.get("discounted_price"),
                    "metal_area_m2": item.get("metal_area_m2"),
                    "weight_kg": item.get("weight_kg"),
                    "notes": item.get("notes"),
                    "project_id": new_id,
                }
                ProductRepository.create(pdata)
                copied += 1
            log_action(
                "project.duplicate",
                entity_type="project",
                entity_id=new_id,
                details={"source_project_id": project_id, "products_copied": copied},
                actor=self._audit_actor(),
            )
            self._load_data()
            if self.main_window:
                self.main_window.set_active_project(new_id)
            QMessageBox.information(
                self, "Успіх", f"Створено копію проєкту ID {new_id}. Скопійовано виробів: {copied}."
            )
        except Exception as e:
            QMessageBox.critical(self, "Помилка", f"Не вдалося дублювати проєкт: {e}")

    def _on_new_project(self):
        dlg = ProjectEditDialog(parent=self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            data = dlg.get_data()
            if not data["name"]:
                QMessageBox.warning(self, "Помилка", "Введіть назву проєкту")
                return
            if not data.get("project_number"):
                data["project_number"] = f"PRJ-{datetime.now().strftime('%Y%m%d-%H%M%S')}"
            data["created_at"] = datetime.now()
            try:
                created = ProjectRepository.create(data)
                project_id = created["id"]
                log_action(
                    "project.create",
                    entity_type="project",
                    entity_id=project_id,
                    details=data,
                    actor=self._audit_actor(),
                )
                self._load_data()
                if self.main_window:
                    self.main_window.set_active_project(project_id)
                QMessageBox.information(self, "Успіх", f"Проєкт створено (ID: {project_id})")
                self._notify_telegram_project_created({**data, "id": project_id})
            except Exception as e:
                QMessageBox.critical(self, "Помилка", f"Не вдалося створити: {e}")

    def _notify_telegram_project_created(self, project: dict):
        """Фонова відправка звіту про новий проєкт у Telegram (не блокує GUI)."""
        worker = FunctionWorker(notify_project_created, project)
        worker.result.connect(self._on_project_notify_result)
        worker.error.connect(lambda _msg: None)  # сповіщення не критичне
        self._tg_notify_worker = worker  # утримуємо посилання, щоб GC не зібрав
        worker.finished.connect(lambda: setattr(self, "_tg_notify_worker", None))
        worker.start()

    @staticmethod
    def _on_project_notify_result(ok: bool):
        """Результат відправки звіту — лише лог, без діалогів (щоб не заважати)."""
        logger = get_logger("projects_tab")
        if ok:
            logger.info("Telegram: звіт про новий проєкт надіслано")
        else:
            logger.warning("Telegram: звіт про новий проєкт НЕ надіслано (бот чи мережа)")

    def _on_edit(self):
        project_id = self._get_selected_id()
        if not project_id:
            QMessageBox.warning(self, "Увага", "Виберіть проєкт для редагування")
            return
        project_data = next((p for p in self._projects if p["id"] == project_id), None)
        if not project_data:
            return
        dlg = ProjectEditDialog(project_data, parent=self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            data = dlg.get_data()
            if not data["name"]:
                QMessageBox.warning(self, "Помилка", "Введіть назву проєкту")
                return
            try:
                ProjectRepository.update(project_id, data)
                log_action(
                    "project.update",
                    entity_type="project",
                    entity_id=project_id,
                    details=data,
                    actor=self._audit_actor(),
                )
                self._load_data()
                QMessageBox.information(self, "Успіх", "Проєкт оновлено!")
                old_status = project_data.get("status") or ""
                if old_status != (data.get("status") or ""):
                    self._notify_status_telegram({**data, "id": project_id}, old_status)
            except Exception as e:
                QMessageBox.critical(self, "Помилка", f"Не вдалося оновити: {e}")

    def _notify_status_telegram(self, project: dict, old_status: str):
        """Фонове сповіщення про зміну статусу проєкту (не критичне)."""
        worker = FunctionWorker(
            notify_project_status_changed, project, old_status, project.get("status") or ""
        )
        worker.error.connect(lambda _msg: None)
        self._tg_status_worker = worker
        worker.finished.connect(lambda: setattr(self, "_tg_status_worker", None))
        worker.start()

    def _on_delete(self):
        project_id = self._get_selected_id()
        if not project_id:
            QMessageBox.warning(self, "Увага", "Виберіть проєкт для видалення")
            return
        project_name = next((p["name"] for p in self._projects if p["id"] == project_id), "")
        msg = 'Видалити проєкт "' + project_name + '" (ID: ' + str(project_id) + ")?"
        msg += " ВСІ вироби та документи цього проєкту також будуть видалені!"
        reply = QMessageBox.question(
            self, "Видалення", msg, QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if reply == QMessageBox.StandardButton.Yes:
            try:
                ProjectRepository.delete_cascade(project_id)
                log_action(
                    "project.delete",
                    entity_type="project",
                    entity_id=project_id,
                    details={"name": project_name},
                    actor=self._audit_actor(),
                )
                self._load_data()
                QMessageBox.information(self, "Успіх", "Проєкт, вироби та документи видалено!")
            except Exception as e:
                QMessageBox.critical(self, "Помилка", f"Не вдалося видалити: {e}")
