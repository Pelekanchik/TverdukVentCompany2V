"""Тести хвилі v2.9: прибутковість, монтажі, склад, історія цін, пошук, прострочення."""

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

pytest.importorskip("PySide6")

import shiboken6
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication

from ventilation_company.database.base import Base
from ventilation_company.services.project_profit import compute_financials
from ventilation_company.services.receivables import is_overdue
from ventilation_company.services.schedule_service import (
    in_period,
    list_crews,
    list_scheduled_works,
    period_bounds,
)


@pytest.fixture(scope="module")
def qapp():
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


@pytest.fixture(autouse=True)
def _destroy_toplevel_widgets(qapp):
    """Детерміновано знищити всі top-level віджети після кожного тесту.

    Без цього непоказані діалоги (WorkEditDialog, вкладки) доживають до
    завершення інтерпретатора і PySide6 падає з access violation при
    розбиранні Qt (процес повертає ненульовий код попри «N passed»).
    """
    yield
    for widget in qapp.topLevelWidgets():
        if shiboken6.isValid(widget):
            widget.hide()
            shiboken6.delete(widget)
    qapp.processEvents()


# ── Фінансовий зріз проєкту ──────────────────────────────────────────────


class TestProjectProfit:
    def test_profit_and_margin(self):
        products = [
            {"total_price": 10000, "discounted_price": 0, "cost_price": 6000, "quantity": 1},
        ]
        works = [{"total_price": 2000}]
        expenses = [
            {"total_price": 500, "direction": "minus"},
            {"total_price": 300, "direction": "plus"},
        ]
        payments = [{"amount": 8000, "type": "вхідний"}]
        f = compute_financials(products, works, expenses, payments)
        assert f["income"] == 12300.0  # 10000 + 2000 + 300
        assert f["cost"] == 6500.0  # 6000 + 500
        assert f["profit"] == 5800.0
        assert f["margin_pct"] == round(5800 / 12300 * 100, 1)
        assert f["paid"] == 8000.0
        assert f["balance"] == 4300.0

    def test_project_discount_overrides_products(self):
        products = [
            {"total_price": 10000, "discounted_price": 0, "cost_price": 5000, "quantity": 1}
        ]
        f = compute_financials(products, [], [], [], discounted_price=9000)
        assert f["income"] == 9000.0

    def test_empty_project(self):
        f = compute_financials([], [], [], [])
        assert f["income"] == 0.0
        assert f["margin_pct"] == 0.0
        assert f["balance"] == 0.0

    def test_outgoing_payment_reduces_paid(self):
        payments = [
            {"amount": 10000, "type": "вхідний"},
            {"amount": 1000, "type": "вихідний"},
        ]
        f = compute_financials([], [], [], payments)
        assert f["paid"] == 9000.0


# ── Планування (чисті функції) ───────────────────────────────────────────


class TestScheduleService:
    def test_period_bounds_today(self):
        from datetime import date

        today = date(2026, 10, 5)
        assert period_bounds("today", today) == (today, today)

    def test_period_bounds_week(self):
        from datetime import date

        start, end = period_bounds("week", date(2026, 10, 7))  # середа
        assert start == date(2026, 10, 5)  # понеділок
        assert end == date(2026, 10, 11)

    def test_period_bounds_month(self):
        from datetime import date

        start, end = period_bounds("month", date(2026, 10, 15))
        assert start == date(2026, 10, 1)
        assert end == date(2026, 10, 31)

    def test_period_bounds_all(self):
        assert period_bounds("all") == (None, None)

    def test_in_period(self):
        from datetime import date

        assert in_period("2026-10-05", None, None)
        assert in_period("2026-10-05", date(2026, 10, 1), date(2026, 10, 31))
        assert not in_period("2026-09-30", date(2026, 10, 1), date(2026, 10, 31))
        assert not in_period("", None, None)
        assert not in_period("невалідна", None, None)

    def test_list_scheduled_works_filters(self, monkeypatch):
        projects = [{"id": 1, "project_number": "ПР-1", "name": "А", "status": "в роботі"}]
        works = [
            {
                "id": 1,
                "project_id": 1,
                "work_name": "Монтаж",
                "work_date": "2026-10-05",
                "crew": "Бригада №1",
                "total_price": 3000.0,
            },
            {
                "id": 2,
                "project_id": 1,
                "work_name": "Доставка",
                "work_date": "",
                "crew": "",
                "total_price": 500.0,
            },
        ]
        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectRepository.list_all",
            staticmethod(lambda: projects),
        )
        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectWorkRepository.get_all",
            staticmethod(lambda pid: works),
        )
        rows = list_scheduled_works(crew="", period="all")
        assert len(rows) == 1  # без дати не потрапляє
        assert rows[0]["work_name"] == "Монтаж"
        assert rows[0]["crew"] == "Бригада №1"

        assert list_scheduled_works(crew="бригада №1") == rows  # case-insensitive
        assert list_scheduled_works(crew="бригада №2") == []

    def test_list_crews_unique_sorted(self, monkeypatch):
        projects = [{"id": 1, "name": "А", "project_number": "ПР-1"}]
        works = [
            {"crew": "Бригада №2", "work_date": "2026-10-01"},
            {"crew": "Бригада №1", "work_date": "2026-10-02"},
            {"crew": "Бригада №2", "work_date": "2026-10-03"},
            {"crew": "", "work_date": ""},
        ]
        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectRepository.list_all",
            staticmethod(lambda: projects),
        )
        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectWorkRepository.get_all",
            staticmethod(lambda pid: works),
        )
        assert list_crews() == ["Бригада №1", "Бригада №2"]

    def test_list_scheduled_works_done_filter(self, monkeypatch):
        projects = [{"id": 1, "project_number": "ПР-1", "name": "А"}]
        works = [
            {"id": 1, "work_name": "Монтаж", "work_date": "2026-10-05", "is_done": False},
            {"id": 2, "work_name": "Доставка", "work_date": "2026-10-06", "is_done": True},
        ]
        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectRepository.list_all",
            staticmethod(lambda: projects),
        )
        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectWorkRepository.get_all",
            staticmethod(lambda pid: works),
        )
        active = list_scheduled_works(done="active")
        assert [r["work_name"] for r in active] == ["Монтаж"]
        assert active[0]["is_done"] is False

        done = list_scheduled_works(done="done")
        assert [r["work_name"] for r in done] == ["Доставка"]
        assert done[0]["is_done"] is True

        all_rows = list_scheduled_works(done="all")
        assert len(all_rows) == 2

    def test_schedule_tab_toggle_done(self, qapp, monkeypatch):
        """Галочка «виконано» у таблиці одразу зберігається в БД."""
        from ventilation_company.gui_pyside6.schedule_tab import ScheduleTab

        works = [
            {
                "id": 7,
                "project_id": 1,
                "work_name": "Монтаж",
                "work_date": "2026-10-05",
                "crew": "Бригада №1",
                "total_price": 1000.0,
                "is_done": False,
            }
        ]
        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectRepository.list_all",
            staticmethod(lambda: [{"id": 1, "name": "А", "project_number": "ПР-1"}]),
        )
        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectWorkRepository.get_all",
            staticmethod(lambda pid: works),
        )
        updates = []

        class _Repo:
            @staticmethod
            def update(work_id, data):
                updates.append((work_id, data))
                return True

        monkeypatch.setattr(
            "ventilation_company.gui_pyside6.schedule_tab.ProjectWorkRepository",
            _Repo,
        )
        tab = ScheduleTab()
        tab.combo_done.setCurrentIndex(2)  # «Усі» — рядок лишається після позначки
        assert tab.table.rowCount() == 1
        check = tab.table.item(0, 0)
        assert check is not None and check.checkState() == Qt.CheckState.Unchecked
        check.setCheckState(Qt.CheckState.Checked)
        assert updates == [(7, {"is_done": True})]


# ── Прострочена заборгованість ───────────────────────────────────────────


class TestOverdue:
    def test_finished_with_debt(self):
        assert is_overdue("Завершено", 1000)
        assert is_overdue("completed", 0.01)

    def test_finished_without_debt(self):
        assert not is_overdue("Завершено", 0)
        assert not is_overdue("Завершено", -500)

    def test_active_project_never_overdue(self):
        assert not is_overdue("В роботі", 5000)
        assert not is_overdue("", 5000)


# ── Репозиторії (in-memory SQLite) ───────────────────────────────────────


@pytest.fixture()
def wh_repo(monkeypatch):
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()

    class _Ctx:
        def __enter__(self):
            return session

        def __exit__(self, *args):
            return False

    from ventilation_company.database.repositories import warehouse_repo as repo_module

    monkeypatch.setattr(repo_module, "get_db", lambda: _Ctx())
    yield repo_module.WarehouseRepository
    session.close()


@pytest.fixture()
def pp_repo(monkeypatch):
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()

    class _Ctx:
        def __enter__(self):
            return session

        def __exit__(self, *args):
            return False

    from ventilation_company.database.repositories import purchase_price_repo as repo_module

    monkeypatch.setattr(repo_module, "get_db", lambda: _Ctx())
    yield repo_module.PurchasePriceRepository
    session.close()


@pytest.fixture()
def work_repo(monkeypatch):
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()

    class _Ctx:
        def __enter__(self):
            return session

        def __exit__(self, *args):
            return False

    from ventilation_company.database.repositories import project_work_repo as repo_module

    monkeypatch.setattr(repo_module, "get_db", lambda: _Ctx())
    yield repo_module.ProjectWorkRepository
    session.close()


class TestWarehouseRepository:
    def test_create_and_list(self, wh_repo):
        wh_repo.create_item(name="Саморізи", unit="шт", min_quantity=100)
        items = wh_repo.list_items()
        assert len(items) == 1
        assert items[0]["name"] == "Саморізи"
        assert items[0]["quantity"] == 0
        assert items[0]["low"]  # 0 ≤ 100 → на межі

    def test_move_in_updates_quantity(self, wh_repo):
        item = wh_repo.create_item(name="Стрічка", unit="м")
        wh_repo.add_move(item_id=item["id"], quantity=50, kind="in")
        items = wh_repo.list_items()
        assert items[0]["quantity"] == 50

    def test_move_out_and_low_flag(self, wh_repo):
        item = wh_repo.create_item(name="Герметик", unit="шт", min_quantity=10)
        wh_repo.add_move(item_id=item["id"], quantity=10, kind="in")
        wh_repo.add_move(item_id=item["id"], quantity=5, kind="out")
        items = wh_repo.list_items()
        assert items[0]["quantity"] == 5
        assert items[0]["low"]  # 5 ≤ 10
        assert len(wh_repo.low_stock()) == 1

    def test_move_rejects_nonpositive(self, wh_repo):
        item = wh_repo.create_item(name="X")
        with pytest.raises(ValueError):
            wh_repo.add_move(item_id=item["id"], quantity=0, kind="in")

    def test_move_unknown_item_returns_none(self, wh_repo):
        assert wh_repo.add_move(item_id=999, quantity=1, kind="in") is None

    def test_list_moves_ordered(self, wh_repo):
        item = wh_repo.create_item(name="Y")
        wh_repo.add_move(item_id=item["id"], quantity=5, kind="in", move_date="2026-09-01")
        wh_repo.add_move(item_id=item["id"], quantity=2, kind="out", move_date="2026-10-01")
        moves = wh_repo.list_moves(item_id=item["id"])
        assert len(moves) == 2
        assert moves[0]["move_date"] == "2026-10-01"  # новіші першими
        assert moves[0]["kind"] == "out"
        assert moves[1]["item_name"] == "Y"

    def test_delete_item_cascades_moves(self, wh_repo):
        item = wh_repo.create_item(name="Z")
        wh_repo.add_move(item_id=item["id"], quantity=1, kind="in")
        assert wh_repo.delete_item(item["id"])
        assert wh_repo.list_items() == []
        assert wh_repo.list_moves() == []


class TestPurchasePriceRepository:
    def test_record_and_latest(self, pp_repo):
        pp_repo.record(item_name="Лист оцинк. 0.7", price=320, supplier="Метал-База")
        pp_repo.record(item_name="Лист оцинк. 0.7", price=310, supplier="Сталь-Плюс")
        history = pp_repo.latest_for("Лист оцинк. 0.7")
        assert len(history) == 2
        assert history[0]["price"] == 310  # остання першою

    def test_latest_for_unknown_empty(self, pp_repo):
        assert pp_repo.latest_for("Неіснуючий") == []
        assert pp_repo.latest_for("") == []

    def test_best_price(self, pp_repo):
        pp_repo.record(item_name="Вентиль", price=450)
        pp_repo.record(item_name="Вентиль", price=420)
        pp_repo.record(item_name="Вентиль", price=480)
        assert pp_repo.best_price_for("Вентиль") == 420
        assert pp_repo.best_price_for("Інше") == 0.0


class TestProjectWorkRepoDates:
    def test_create_with_date_and_crew(self, work_repo):
        w = work_repo.create(
            {
                "project_id": 1,
                "work_name": "Монтаж",
                "quantity": 2,
                "unit_price": 500,
                "work_date": "2026-10-10",
                "crew": "Бригада №1",
            }
        )
        assert w["work_date"] == "2026-10-10"
        assert w["crew"] == "Бригада №1"
        assert w["total_price"] == 1000

    def test_create_without_date(self, work_repo):
        w = work_repo.create({"project_id": 1, "work_name": "Робота"})
        assert w["work_date"] == ""
        assert w["crew"] == ""

    def test_update_date(self, work_repo):
        w = work_repo.create({"project_id": 1, "work_name": "Робота", "unit_price": 100})
        assert work_repo.update(w["id"], {"work_date": "2026-11-01", "crew": "Б-2"})
        updated = work_repo.get_all(1)[0]
        assert updated["work_date"] == "2026-11-01"
        assert updated["crew"] == "Б-2"

    def test_clear_date_via_empty_string(self, work_repo):
        w = work_repo.create(
            {"project_id": 1, "work_name": "Р", "work_date": "2026-10-01", "crew": "Б-1"}
        )
        work_repo.update(w["id"], {"work_date": ""})
        assert work_repo.get_all(1)[0]["work_date"] == ""


# ── Глобальний пошук (з підміною репозиторіїв) ───────────────────────────


class TestGlobalSearch:
    def test_search_projects_and_products(self, monkeypatch):
        import ventilation_company.services.global_search as gs

        monkeypatch.setattr(
            gs.ProjectRepository,
            "list_all",
            staticmethod(
                lambda: [
                    {
                        "id": 1,
                        "name": "Вентиляція кафе",
                        "project_number": "ПР-7",
                        "client": "ТОВ А",
                    },
                    {"id": 2, "name": "Склад", "project_number": "ПР-8", "client": "ПП Б"},
                ]
            ),
        )
        monkeypatch.setattr(
            gs.ProductRepository,
            "search",
            staticmethod(
                lambda query="", **kw: (
                    [
                        {
                            "project_id": 1,
                            "name": f"Повітропровід {query}",
                            "product_type": "прямокутний",
                        }
                    ]
                    if query
                    else []
                )
            ),
        )
        monkeypatch.setattr(gs.ProjectWorkRepository, "get_all", staticmethod(lambda pid: []))
        monkeypatch.setattr(
            gs.ProjectDrawingRepository, "get_by_project", staticmethod(lambda pid: [])
        )
        monkeypatch.setattr(gs.ClientRepository, "list_all", staticmethod(lambda: []))

        results = gs.search_all("вентиляція")
        kinds = [r["kind"] for r in results]
        assert "projects" in kinds
        project = next(r for r in results if r["kind"] == "projects")
        assert project["project_id"] == 1

        results2 = gs.search_all("повітропровід")
        assert any(r["kind"] == "products" for r in results2)

    def test_empty_query(self):
        import ventilation_company.services.global_search as gs

        assert gs.search_all("") == []
        assert gs.search_all("   ") == []


# ── Excel-експорт КП ─────────────────────────────────────────────────────


class TestProposalExcel:
    def test_export_creates_file(self, tmp_path):
        from openpyxl import load_workbook

        from ventilation_company.proposal_generator import export_proposal_to_excel

        path = str(tmp_path / "кп.xlsx")
        items = [
            {
                "name": "Вентилятор",
                "description": "відцентровий",
                "quantity": 2,
                "unit": "шт",
                "price": 5000,
            },
            {"name": "Монтаж", "description": "роботи", "quantity": 1, "unit": "шт", "price": 3000},
        ]
        project_data = {
            "name": "Кафе",
            "project_number": "ПР-1",
            "client": "ТОВ «Смак»",
            "company": {"name": "ВАТ", "phone": "044-000-00-00"},
        }
        export_proposal_to_excel(project_data, items, path)
        wb = load_workbook(path)
        ws = wb.active
        assert ws["A1"].value == "КОМЕРЦІЙНА ПРОПОЗИЦІЯ"
        # Разом: 2*5000 + 1*3000 = 13000
        total_row = 7 + len(items) + 1
        assert ws.cell(row=total_row, column=7).value == 13000


# ── GUI: нові вкладки та діалоги (без модальних діалогів) ────────────────


class TestNewTabsGui:
    def test_schedule_tab_fill(self, qapp, monkeypatch):
        from ventilation_company.gui_pyside6.schedule_tab import ScheduleTab

        works = [
            {
                "id": 1,
                "project_id": 1,
                "work_name": "Монтаж",
                "work_date": "2026-10-05",
                "crew": "Бригада №1",
                "total_price": 1000.0,
            }
        ]
        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectRepository.list_all",
            staticmethod(lambda: [{"id": 1, "name": "А", "project_number": "ПР-1"}]),
        )
        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectWorkRepository.get_all",
            staticmethod(lambda pid: works),
        )
        tab = ScheduleTab()
        assert tab.table.rowCount() == 1
        assert tab.table.item(0, 3).text() == "Монтаж"  # 0=✓, 1=Дата, 2=Проєкт
        assert tab.table.item(0, 6) is None  # 6 колонок — індекси 0..5
        assert tab.table.item(0, 4).text() == "Бригада №1"

    def test_schedule_tab_empty_on_error(self, qapp, monkeypatch):
        from ventilation_company.gui_pyside6.schedule_tab import ScheduleTab

        def _boom():
            raise RuntimeError("no db")

        monkeypatch.setattr(
            "ventilation_company.services.schedule_service.ProjectRepository.list_all", _boom
        )
        tab = ScheduleTab()
        assert tab.table.rowCount() == 0

    def test_warehouse_tab_fill_and_low(self, qapp, monkeypatch):
        from ventilation_company.gui_pyside6.warehouse_tab import WarehouseTab

        rows = [
            {
                "id": 1,
                "name": "Герметик",
                "unit": "шт",
                "quantity": 3,
                "min_quantity": 10,
                "low": True,
            },
            {
                "id": 2,
                "name": "Стрічка",
                "unit": "м",
                "quantity": 50,
                "min_quantity": 10,
                "low": False,
            },
        ]
        monkeypatch.setattr(
            "ventilation_company.gui_pyside6.warehouse_tab.WarehouseRepository.list_items",
            staticmethod(lambda: rows),
        )
        tab = WarehouseTab()
        assert tab.table.rowCount() == 2
        assert "⚠️" in tab.table.item(0, 4).text()
        assert "✓" in tab.table.item(1, 4).text()
        assert "на межі: 1" in tab.lbl_total.text()

    def test_global_search_dialog(self, qapp, monkeypatch):
        import ventilation_company.services.global_search as gs
        from ventilation_company.gui_pyside6.global_search_dialog import GlobalSearchDialog

        monkeypatch.setattr(
            gs.ProjectRepository,
            "list_all",
            staticmethod(
                lambda: [
                    {
                        "id": 5,
                        "name": "Вентиляція ресторану",
                        "project_number": "ПР-9",
                        "client": "ТОВ В",
                    }
                ]
            ),
        )
        monkeypatch.setattr(gs.ProductRepository, "search", staticmethod(lambda **kw: []))
        monkeypatch.setattr(gs.ProjectWorkRepository, "get_all", staticmethod(lambda pid: []))
        monkeypatch.setattr(
            gs.ProjectDrawingRepository, "get_by_project", staticmethod(lambda pid: [])
        )
        monkeypatch.setattr(gs.ClientRepository, "list_all", staticmethod(lambda: []))

        dlg = GlobalSearchDialog()
        dlg.edit.setText("ресторану")
        assert dlg.list.count() == 1
        assert dlg._results[0]["project_id"] == 5
        dlg.edit.setText("x")  # менше 2 символів → порожньо
        assert dlg.list.count() == 0

    def test_money_tab_overdue_filter(self, qapp, monkeypatch):
        from ventilation_company.gui_pyside6.money_tab import MoneyTab

        rows = [
            {
                "project_id": 1,
                "project_number": "ПР-1",
                "name": "Завершений",
                "client": "А",
                "status": "Завершено",
                "total": 10000.0,
                "paid": 4000.0,
                "balance": 6000.0,
                "percent": 40.0,
                "overpaid": False,
                "payments_count": 1,
                "last_payment": "2026-09-01",
            },
            {
                "project_id": 2,
                "project_number": "ПР-2",
                "name": "Активний",
                "client": "Б",
                "status": "В роботі",
                "total": 5000.0,
                "paid": 1000.0,
                "balance": 4000.0,
                "percent": 20.0,
                "overpaid": False,
                "payments_count": 1,
                "last_payment": "2026-09-15",
            },
        ]
        monkeypatch.setattr(
            "ventilation_company.gui_pyside6.money_tab.build_receivables",
            lambda: rows,
        )
        tab = MoneyTab()
        tab.combo_filter.setCurrentText("Прострочено")
        assert tab.table.rowCount() == 1
        assert tab.table.item(0, 1).text() == "Завершений"
        assert "⏰" in tab.table.item(0, 6).text()

    def test_work_edit_dialog_date(self, qapp):
        from PySide6.QtCore import QDate

        from ventilation_company.gui_pyside6.project_card_dialog import WorkEditDialog

        dlg = WorkEditDialog(project_id=1)
        assert dlg.chk_has_date.isChecked()  # нова робота — з датою за замовчуванням
        dlg.date_edit.setDate(QDate(2026, 10, 20))
        dlg.edit_crew.setText("Бригада №3")
        data = dlg.get_data()
        assert data["work_date"] == "2026-10-20"
        assert data["crew"] == "Бригада №3"

    def test_work_edit_dialog_without_date(self, qapp):
        from ventilation_company.gui_pyside6.project_card_dialog import WorkEditDialog

        dlg = WorkEditDialog(project_id=1)
        dlg.chk_has_date.setChecked(False)
        assert dlg.get_data()["work_date"] == ""

    def test_work_edit_dialog_prefills_from_data(self, qapp):
        from ventilation_company.gui_pyside6.project_card_dialog import WorkEditDialog

        dlg = WorkEditDialog(
            project_id=1,
            work_data={
                "id": 1,
                "work_name": "Монтаж",
                "work_date": "2026-09-15",
                "crew": "Бригада №1",
            },
        )
        assert dlg.chk_has_date.isChecked()
        assert dlg.date_edit.date().toString("yyyy-MM-dd") == "2026-09-15"
        assert dlg.edit_crew.text() == "Бригада №1"
        assert dlg.get_data()["work_date"] == "2026-09-15"

    def test_quick_work_dialog_create(self, qapp, monkeypatch):
        from PySide6.QtCore import QDate

        from ventilation_company.gui_pyside6.schedule_tab import QuickWorkDialog

        monkeypatch.setattr(
            "ventilation_company.gui_pyside6.schedule_tab.ProjectRepository.list_all",
            staticmethod(
                lambda: [{"id": 7, "name": "Кафе", "project_number": "ПР-7", "status": "в роботі"}]
            ),
        )
        created = []

        class _FakeRepo:
            @staticmethod
            def create(data):
                created.append(data)
                return data

        monkeypatch.setattr(
            "ventilation_company.gui_pyside6.schedule_tab.ProjectWorkRepository", _FakeRepo
        )
        dlg = QuickWorkDialog()
        assert dlg.combo_project.currentData() == 7
        dlg.edit_name.setText("Монтаж вентилятора")
        dlg.date_edit.setDate(QDate(2026, 10, 25))
        dlg.edit_crew.setText("Бригада №2")
        dlg.spin_price.setValue(1500)
        assert dlg.create_work()
        assert created[0]["project_id"] == 7
        assert created[0]["work_date"] == "2026-10-25"
        assert created[0]["crew"] == "Бригада №2"

    def test_quick_work_dialog_requires_name(self, qapp, monkeypatch):
        from ventilation_company.gui_pyside6.schedule_tab import QuickWorkDialog

        monkeypatch.setattr(
            "ventilation_company.gui_pyside6.schedule_tab.ProjectRepository.list_all",
            staticmethod(lambda: [{"id": 1, "name": "А", "project_number": "ПР-1"}]),
        )
        monkeypatch.setattr(
            "ventilation_company.gui_pyside6.schedule_tab.QMessageBox.warning",
            staticmethod(lambda *a, **k: None),
        )
        dlg = QuickWorkDialog()
        dlg.edit_name.setText("")
        assert not dlg.create_work()
