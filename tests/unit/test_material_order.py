"""Тести заявки на матеріали: калькулятор, Excel-експортер, кнопка у картці проєкту."""

import time

import pytest

pytest.importorskip("PySide6")

from PySide6.QtWidgets import QApplication, QLineEdit, QMessageBox

from ventilation_company.material_order import (
    MaterialOrderExporter,
    calculate_material_order,
)

PRODUCTS = [
    {
        "name": "Повітропровід",
        "product_type": "повітропровід прямокутний",
        "material": "оцинкована сталь",
        "thickness": 0.7,
        "width": 400,
        "height": 200,
        "length": 1000,
        "quantity": 10,
        "metal_area_m2": 1.2,
        "has_flanges": True,
        "flange_count": 2,
        "components": ["Вентилятор"],
    }
]


class _PricingStub:
    material_prices: dict = {}

    def reload(self):
        return None

    def get_material_price(self, material, thickness, default=0):
        """Як у production: нечутливо до регістру назви та формату товщини."""
        wanted = str(material or "").strip().lower()
        wanted_th = f"{float(thickness):g}"
        for mat_name, thicknesses in self.material_prices.items():
            if not isinstance(thicknesses, dict):
                continue
            if str(mat_name).strip().lower() != wanted:
                continue
            for th_key, price in thicknesses.items():
                if f"{float(th_key):g}" == wanted_th:
                    return float(price)
        return float(default)


class _BusinessStub:
    def get_extra_material_price(self, key, default=0):
        return default

    def get_component(self, key):
        return {}


class _BusinessStringStub(_BusinessStub):
    """Імітує реальний файл, де ціни збережені рядками («180»)."""

    def get_extra_material_price(self, key, default=0):
        return "180"

    def get_component(self, key):
        return {"ціна": "1 598,15", "одиниця": "шт"}


@pytest.fixture(autouse=True)
def _stub_settings(monkeypatch):
    """Не читаємо реальні файли налаштувань у юніт-тестах."""
    monkeypatch.setattr(
        "ventilation_company.services.pricing_settings.PricingSettings.get_instance",
        staticmethod(lambda: _PricingStub()),
    )
    monkeypatch.setattr(
        "ventilation_company.services.business_settings.BusinessSettings.get_instance",
        staticmethod(lambda: _BusinessStub()),
    )


@pytest.fixture(scope="module")
def qapp():
    """Власна фікстура QApplication (pytest-qt у CI не встановлено)."""
    app = QApplication.instance() or QApplication([])
    yield app


class TestMaterialCalculator:
    def test_calculator_covers_all_categories(self):
        order = calculate_material_order(PRODUCTS, project_name="Тест")
        cats = {i.category for i in order.items}
        assert "Листовий метал" in cats
        assert "Ущільнювачі" in cats
        assert "Кріплення" in cats
        assert "Ізоляція" in cats
        assert "Комплектуючі" in cats
        assert "Розхідні матеріали" in cats
        assert order.project_name == "Тест"

    def test_bolt_count_medium_flange(self):
        # max_dim = 400 → 6 болтів на фланець; 2 фланці × 10 виробів = 20 фланців → 120
        order = calculate_material_order(PRODUCTS, project_name="Тест")
        bolts = [i for i in order.get_by_category("Кріплення") if i.name.startswith("Болт")]
        assert bolts and bolts[0].quantity == 120

    def test_empty_products_gives_only_consumables(self):
        order = calculate_material_order([], project_name="Порожній")
        assert order.items
        assert all(i.category == "Розхідні матеріали" for i in order.items)

    def test_string_prices_do_not_break_total(self, monkeypatch):
        """Регресія: ціни-рядки з business_settings.json («180») не дають TypeError."""
        monkeypatch.setattr(
            "ventilation_company.services.business_settings.BusinessSettings.get_instance",
            staticmethod(lambda: _BusinessStringStub()),
        )
        order = calculate_material_order(PRODUCTS, project_name="Тест")
        assert isinstance(order.total_cost, float)
        wool = [i for i in order.get_by_category("Ізоляція") if "вата" in i.name]
        assert wool and wool[0].price_per_unit == 180.0
        comp = order.get_by_category("Комплектуючі")
        assert comp and comp[0].price_per_unit == 1598.15

    def test_capitalized_material_gets_price(self, monkeypatch):
        """Регресія: «Оцинкована сталь» (з великої літери, як у виробі) знаходить
        ціну «оцинкована сталь» з ціноутворення — раніше ціна була 0."""
        pricing = _PricingStub()
        pricing.material_prices = {"оцинкована сталь": {"0.7": 580.0}}
        monkeypatch.setattr(
            "ventilation_company.services.pricing_settings.PricingSettings.get_instance",
            staticmethod(lambda: pricing),
        )
        products = [
            {
                "name": "Повітропровід",
                "product_type": "повітропровід круглий",
                "material": "Оцинкована сталь",
                "thickness": 0.7,
                "width": 400,
                "height": 0,
                "length": 1000,
                "quantity": 10,
                "metal_area_m2": 12.0,
            }
        ]
        order = calculate_material_order(products, project_name="Тест")
        metal = order.get_by_category("Листовий метал")
        assert metal, "рядок металу має бути"
        # 580 ₴/м² × 3.125 м² (лист 1250×2500) = 1812.50 ₴/лист
        assert metal[0].price_per_unit == 1812.50


class TestMaterialOrderExporter:
    def test_export_creates_excel(self, tmp_path):
        from openpyxl import load_workbook

        order = calculate_material_order(PRODUCTS, project_name="Тест")
        path = tmp_path / "заявка.xlsx"
        MaterialOrderExporter(order).export(str(path))
        assert path.exists()
        wb = load_workbook(path)
        ws = wb.active
        assert "ЗАЯВКА НА МАТЕРІАЛИ" in ws["A1"].value
        assert "Тест" in ws["A1"].value


CARD_DATA = {
    "project": {
        "id": 1,
        "name": "Тестовий проєкт",
        "project_number": "PRJ-2026-001",
        "client": "ТОВ «Клієнт»",
        "status": "в роботі",
        "created_at": "2026-09-27 10:00",
        "cost_price": 5000.0,
        "customer_price": 8000.0,
        "discounted_price": 0,
        "works_total": 0.0,
        "plus_expenses_total": 0.0,
        "minus_expenses_total": 0.0,
        "paid_total": 0.0,
    },
    "products": PRODUCTS,
    "documents": [],
    "drawings": [],
    "works": [],
    "expenses": [],
    "payments": [],
}


def _wait_worker(qapp, dlg, iterations=500):
    for _ in range(iterations):
        qapp.processEvents()
        if dlg._worker is None:
            return True
        time.sleep(0.001)
    return False


class TestMaterialOrderButton:
    def test_button_saves_excel(self, qapp, monkeypatch, tmp_path):
        from ventilation_company.gui_pyside6.project_card_dialog import ProjectCardDialog

        monkeypatch.setattr(ProjectCardDialog, "_fetch_data", lambda self: CARD_DATA)
        monkeypatch.setattr(
            "PySide6.QtWidgets.QMessageBox.warning", staticmethod(lambda *a, **k: None)
        )
        monkeypatch.setattr(
            "PySide6.QtWidgets.QMessageBox.critical", staticmethod(lambda *a, **k: None)
        )
        monkeypatch.setattr(
            "PySide6.QtWidgets.QMessageBox.information", staticmethod(lambda *a, **k: None)
        )
        monkeypatch.setattr(
            "PySide6.QtWidgets.QMessageBox.question",
            staticmethod(lambda *a, **k: QMessageBox.StandardButton.No),
        )
        out = tmp_path / "заявка.xlsx"
        monkeypatch.setattr(
            "PySide6.QtWidgets.QFileDialog.getSaveFileName",
            staticmethod(lambda *a, **k: (str(out), "Excel (*.xlsx)")),
        )
        from PySide6.QtWidgets import QDialog

        monkeypatch.setattr(
            "ventilation_company.gui_pyside6.project_card_dialog.MaterialOrderPreviewDialog.exec",
            lambda self: QDialog.DialogCode.Accepted,
        )

        dlg = ProjectCardDialog(1)
        assert _wait_worker(qapp, dlg), "worker завис"
        dlg._on_material_order()
        assert out.exists()
        dlg.close()

    def test_no_products_shows_info_and_saves_nothing(self, qapp, monkeypatch, tmp_path):
        from ventilation_company.gui_pyside6.project_card_dialog import ProjectCardDialog

        data = dict(CARD_DATA)
        data["products"] = []
        monkeypatch.setattr(ProjectCardDialog, "_fetch_data", lambda self: data)
        monkeypatch.setattr(
            "PySide6.QtWidgets.QMessageBox.warning", staticmethod(lambda *a, **k: None)
        )
        monkeypatch.setattr(
            "PySide6.QtWidgets.QMessageBox.critical", staticmethod(lambda *a, **k: None)
        )
        called = {"info": False}
        monkeypatch.setattr(
            "PySide6.QtWidgets.QMessageBox.information",
            staticmethod(lambda *a, **k: called.__setitem__("info", True)),
        )
        dlg = ProjectCardDialog(1)
        assert _wait_worker(qapp, dlg), "worker завис"
        dlg._on_material_order()
        assert called["info"], "очікувалося вікно «немає виробів»"
        dlg.close()


class TestMaterialOrderPreviewDialog:
    def _make_dialog(self, qapp, monkeypatch):
        from ventilation_company.gui_pyside6.material_order_dialog import (
            MaterialOrderPreviewDialog,
        )
        from ventilation_company.material_order import MaterialItem, MaterialOrder

        monkeypatch.setattr(
            "PySide6.QtWidgets.QMessageBox.information", staticmethod(lambda *a, **k: None)
        )
        monkeypatch.setattr(
            "PySide6.QtWidgets.QMessageBox.warning", staticmethod(lambda *a, **k: None)
        )
        order = MaterialOrder(
            project_name="Тест",
            items=[
                MaterialItem(
                    category="Кріплення",
                    name="Болт М8",
                    specification="DIN 933",
                    unit="шт",
                    quantity=10,
                    price_per_unit=3.5,
                ),
                MaterialItem(
                    category="Ізоляція",
                    name="Мінвата",
                    specification="50 мм",
                    unit="м²",
                    quantity=20,
                    price_per_unit=180,
                ),
            ],
        )
        return MaterialOrderPreviewDialog(order)

    def test_edits_reflected_in_get_order(self, qapp, monkeypatch):
        dlg = self._make_dialog(qapp, monkeypatch)
        assert dlg.table.rowCount() == 2
        # Змінюємо кількість болтів 10 → 25
        dlg.table.item(0, 4).setText("25")
        order = dlg.get_order()
        assert order.items[0].quantity == 25
        assert order.items[0].price_per_unit == 3.5
        # Сума рядка перерахована
        assert dlg.table.item(0, 6).text() == "87.50"
        dlg.close()

    def test_delete_and_add_row(self, qapp, monkeypatch):
        dlg = self._make_dialog(qapp, monkeypatch)
        dlg.table.setCurrentCell(1, 0)
        dlg._delete_selected_row()
        assert dlg.table.rowCount() == 1
        dlg._add_row()
        assert dlg.table.rowCount() == 2
        order = dlg.get_order()
        assert order.items[1].name == "Новий матеріал"
        assert order.items[1].quantity == 1
        dlg.close()

    def test_get_order_handles_bad_numbers(self, qapp, monkeypatch):
        dlg = self._make_dialog(qapp, monkeypatch)
        dlg.table.item(0, 4).setText("abc")  # нечислове → 0
        order = dlg.get_order()
        assert order.items[0].quantity == 0.0
        dlg.close()

    def test_duplicate_selected_row(self, qapp, monkeypatch):
        dlg = self._make_dialog(qapp, monkeypatch)
        assert dlg.table.rowCount() == 2
        dlg.table.setCurrentCell(0, 1)
        dlg._duplicate_selected_row()
        assert dlg.table.rowCount() == 3
        # Копія вставлена одразу після оригіналу
        assert dlg.table.item(1, 1).text() == "Болт М8"
        assert dlg.table.item(1, 4).text() == "10"
        assert dlg.table.item(1, 5).text() == "3.5"
        # Оригінальний другий рядок зсунуто вниз
        assert dlg.table.item(2, 1).text() == "Мінвата"
        order = dlg.get_order()
        assert [i.name for i in order.items] == ["Болт М8", "Болт М8", "Мінвата"]
        dlg.close()

    def test_duplicate_without_selection_warns(self, qapp, monkeypatch):
        from ventilation_company.gui_pyside6.material_order_dialog import (
            MaterialOrderPreviewDialog,
        )
        from ventilation_company.material_order import MaterialOrder

        called = {"info": False}
        monkeypatch.setattr(
            "PySide6.QtWidgets.QMessageBox.information",
            staticmethod(lambda *a, **k: called.update(info=True)),
        )
        order = MaterialOrder(project_name="Тест", items=[])
        dlg = MaterialOrderPreviewDialog(order)
        dlg._duplicate_selected_row()  # currentRow() == -1
        assert called["info"]
        dlg.close()

    def test_cell_editor_fits_row_height(self, qapp, monkeypatch):
        """Редактор комірки поміщається у висоту рядка — текст не обрізається."""
        from ventilation_company.gui_pyside6.theme import Theme

        Theme.apply(qapp)
        dlg = self._make_dialog(qapp, monkeypatch)
        dlg.show()
        qapp.processEvents()
        dlg.table.editItem(dlg.table.item(0, 1))
        qapp.processEvents()
        editors = dlg.table.findChildren(QLineEdit)
        assert editors, "редактор комірки не відкрився"
        row_height = dlg.table.rowHeight(0)
        for editor in editors:
            assert (
                editor.sizeHint().height() <= row_height
            ), f"редактор ({editor.sizeHint().height()}px) вищий за рядок ({row_height}px)"
        dlg.close()
